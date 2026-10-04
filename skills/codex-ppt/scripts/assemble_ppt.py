#!/usr/bin/env python3
"""
PPT 조립 스크립트

디렉터리에 있는 슬라이드 이미지들을 PowerPoint 프레젠테이션으로 조립한다.
각 이미지는 한 페이지 전체를 채우는 슬라이드로 삽입된다.
"""

import argparse
import os
import re
import sys
import tempfile
from typing import Dict, List, Optional, Tuple


def dependency_hint() -> str:
    runtime_home = os.path.expanduser(os.environ.get("CODEX_PPT_HOME", "~/.codex-ppt-skill"))
    python = os.path.join(
        runtime_home,
        ".venv",
        "Scripts" if os.name == "nt" else "bin",
        "python.exe" if os.name == "nt" else "python",
    )
    runtime_script = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "codex_ppt_runtime.py",
    )
    return (
        f"다음을 실행하세요: python3 {runtime_script} bootstrap\n"
        f"또는 직접 실행: {python} -m pip install -r "
        f"{os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'requirements.txt')}"
    )


def get_slide_images(ppt_project_dir: str) -> List[str]:
    """
    슬라이드 이미지 파일 목록을 페이지 번호순으로 가져온다

    Args:
        ppt_project_dir: PPT 프로젝트 디렉터리 (origin_image 하위 디렉터리 포함)

    Returns:
        페이지 순서대로 정렬된 이미지 파일 경로 목록
    """
    # origin_image 하위 디렉터리에서 읽는다
    origin_image_dir = os.path.join(ppt_project_dir, "origin_image")

    if not os.path.exists(origin_image_dir):
        print(f"오류: origin_image 디렉터리가 없습니다: {origin_image_dir}")
        return []

    print(f"origin_image 디렉터리에서 이미지 읽는 중: {origin_image_dir}")

    # 지원하는 이미지 형식
    image_extensions = {'.png', '.jpg', '.jpeg', '.gif', '.bmp'}

    # 정식 슬라이드 이미지만 가져온다. sample_slide.png, 초안, 참고 이미지가 PPT에 잘못 들어가는 것을 막는다.
    slide_name_pattern = re.compile(r"^slide_(\d+)\.(png|jpe?g|gif|bmp)$", re.IGNORECASE)
    image_files = []
    for file in os.listdir(origin_image_dir):
        file_path = os.path.join(origin_image_dir, file)
        if os.path.isfile(file_path):
            ext = os.path.splitext(file)[1].lower()
            if ext in image_extensions and slide_name_pattern.match(file):
                image_files.append(file_path)

    def slide_sort_key(path: str) -> Tuple[int, str]:
        filename = os.path.basename(path)
        stem = os.path.splitext(filename)[0]
        return (int(re.search(r"^slide_(\d+)", stem).group(1)), filename.lower())

    # 페이지 번호순 정렬 — slide_10이 slide_2보다 앞에 오는 것을 방지한다.
    # 같은 페이지 번호에 파일이 여러 개면 파일명으로 안정 정렬한다.
    image_files.sort(key=slide_sort_key)

    return image_files


def load_speaker_notes(ppt_project_dir: str) -> Dict[int, str]:
    """
    speech.md에서 페이지별 발표자 노트를 읽는다.

    지원하는 제목 형식:
    - ## Slide 1: 제목
    - ### Slide 1: 제목
    - ## 第 1 页：제목  (업스트림 레거시 형식 호환용)
    """
    speech_path = os.path.join(ppt_project_dir, "speech.md")
    if not os.path.exists(speech_path):
        return {}

    with open(speech_path, "r", encoding="utf-8") as f:
        content = f.read()

    notes: Dict[int, str] = {}
    current_slide: Optional[int] = None
    current_lines: List[str] = []
    # '第 N 页' 대안은 업스트림에서 작성된 레거시 speech.md 호환용으로 유지한다.
    heading_pattern = re.compile(r"^#{2,4}\s*(?:Slide\s*(\d+)|第\s*(\d+)\s*页)\b.*$", re.IGNORECASE)

    def flush_current() -> None:
        if current_slide is None:
            return
        text = "\n".join(current_lines).strip()
        if text:
            notes[current_slide] = text

    for line in content.splitlines():
        match = heading_pattern.match(line.strip())
        if match:
            flush_current()
            current_slide = int(match.group(1) or match.group(2))
            current_lines = []
            continue

        if current_slide is not None:
            current_lines.append(line)

    flush_current()
    return notes


def compress_image_if_needed(
    image_path: str,
    max_size_mb: float = 2.0,
    quality_step: int = 5
) -> Optional[str]:
    """
    이미지가 지정 크기를 초과하면 압축하고 임시 파일 경로를 반환한다

    Args:
        image_path: 원본 이미지 경로
        max_size_mb: 최대 파일 크기 (MB)
        quality_step: 품질을 단계적으로 낮추는 폭

    Returns:
        str: 압축이 필요하면 임시 파일 경로, 아니면 None
    """
    max_size_bytes = max_size_mb * 1024 * 1024

    # 원본 파일 크기 확인
    file_size = os.path.getsize(image_path)

    if file_size <= max_size_bytes:
        # 압축 불필요
        return None

    print(f"  이미지 크기 {file_size / 1024 / 1024:.2f}MB, 압축 필요...")

    try:
        from PIL import Image

        # 이미지 열기
        img = Image.open(image_path)

        # RGBA를 RGB로 변환 (JPEG 저장이 필요한 경우)
        if img.mode in ('RGBA', 'LA', 'P'):
            background = Image.new('RGB', img.size, (255, 255, 255))
            if img.mode == 'P':
                img = img.convert('RGBA')
            background.paste(img, mask=img.split()[-1] if img.mode in ('RGBA', 'LA') else None)
            img = background

        # 임시 파일 생성
        temp_fd, temp_path = tempfile.mkstemp(suffix='.jpg')
        os.close(temp_fd)

        # 높은 품질부터 단계적으로 압축 시도
        quality = 95
        while quality > 20:
            img.save(temp_path, 'JPEG', quality=quality, optimize=True)
            compressed_size = os.path.getsize(temp_path)

            if compressed_size <= max_size_bytes:
                print(f"  압축 성공: {compressed_size / 1024 / 1024:.2f}MB (품질: {quality})")
                return temp_path

            quality -= quality_step

        # 그래도 크면 크기 축소 시도
        print("  품질 압축으로 부족하여 크기를 줄입니다...")
        scale = 0.9
        while scale > 0.3:
            new_width = int(img.width * scale)
            new_height = int(img.height * scale)
            resized_img = img.resize((new_width, new_height), Image.Resampling.LANCZOS)

            resized_img.save(temp_path, 'JPEG', quality=85, optimize=True)
            compressed_size = os.path.getsize(temp_path)

            if compressed_size <= max_size_bytes:
                print(f"  압축 성공: {compressed_size / 1024 / 1024:.2f}MB (스케일: {scale:.0%})")
                return temp_path

            scale -= 0.1

        # 더 이상 줄일 수 없으면 마지막 결과를 반환
        print(f"  경고: {max_size_mb}MB 이하로 압축할 수 없어 최소 크기 버전을 사용합니다")
        return temp_path

    except ImportError:
        print("오류: Pillow 라이브러리가 설치되어 있지 않습니다")
        print(dependency_hint())
        return None
    except Exception as e:
        print(f"  경고: 이미지 압축 실패: {e}")
        if 'temp_path' in locals() and os.path.exists(temp_path):
            os.remove(temp_path)
        return None


def create_presentation(
    image_files: List[str],
    output_path: str,
    aspect_ratio: str = "16:9",
    speaker_notes: Optional[Dict[int, str]] = None,
) -> bool:
    """
    PowerPoint 프레젠테이션을 생성한다

    Args:
        image_files: 슬라이드 이미지 파일 목록
        output_path: 출력 PPT 파일 경로
        aspect_ratio: 슬라이드 가로세로 비율 (16:9 또는 4:3)

    Returns:
        bool: 성공하면 True, 실패하면 False
    """
    try:
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            print("오류: python-pptx 라이브러리가 설치되어 있지 않습니다")
            print(dependency_hint())
            return False

        # 프레젠테이션 생성
        prs = Presentation()

        # 슬라이드 크기 설정
        if aspect_ratio == "16:9":
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(5.625)
        elif aspect_ratio == "4:3":
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)
        else:
            print(f"경고: 지원하지 않는 가로세로 비율 {aspect_ratio}, 기본값 16:9 사용")
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(5.625)

        speaker_notes = speaker_notes or {}

        # 각 슬라이드 추가
        temp_files_to_cleanup = []
        for i, image_path in enumerate(image_files, 1):
            if not os.path.exists(image_path):
                print(f"경고: 이미지 파일이 없습니다: {image_path}")
                continue

            # 필요하면 이미지 압축
            compressed_path = compress_image_if_needed(image_path, max_size_mb=2.0)

            # 압축된 이미지 또는 원본 사용
            image_to_use = compressed_path if compressed_path else image_path

            # 나중에 정리할 수 있도록 임시 파일 기록
            if compressed_path:
                temp_files_to_cleanup.append(compressed_path)

            # 빈 레이아웃 사용 (인덱스 6)
            blank_slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_slide_layout)

            # 이미지를 페이지 전체를 채우도록 슬라이드에 추가
            slide.shapes.add_picture(
                image_to_use,
                left=0,
                top=0,
                width=prs.slide_width,
                height=prs.slide_height
            )

            note_text = speaker_notes.get(i)
            if note_text:
                notes_frame = slide.notes_slide.notes_text_frame
                notes_frame.clear()
                notes_frame.text = note_text

            print(f"✓ {i} 페이지 추가됨: {os.path.basename(image_path)}")

        # 출력 디렉터리가 없으면 생성
        output_dir = os.path.dirname(output_path)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)

        # 프레젠테이션 저장
        prs.save(output_path)
        print(f"\n✓ PPT 파일이 저장되었습니다: {output_path}")
        print(f"  총 페이지 수: {len(image_files)}")
        if speaker_notes:
            matched_notes = sum(1 for i in range(1, len(image_files) + 1) if speaker_notes.get(i))
            print(f"  노트 기록됨: {matched_notes}/{len(image_files)} 페이지")

        # 임시 파일 정리
        for temp_file in temp_files_to_cleanup:
            try:
                if os.path.exists(temp_file):
                    os.remove(temp_file)
            except Exception as e:
                print(f"경고: 임시 파일 정리 실패 {temp_file}: {e}")

        return True

    except Exception as e:
        print(f"오류: PPT 생성 실패: {e}")
        import traceback
        traceback.print_exc()

        # 실패해도 임시 파일은 정리
        if 'temp_files_to_cleanup' in locals():
            for temp_file in temp_files_to_cleanup:
                try:
                    if os.path.exists(temp_file):
                        os.remove(temp_file)
                except:
                    pass

        return False


def main():
    parser = argparse.ArgumentParser(
        description='슬라이드 이미지들을 PowerPoint 프레젠테이션으로 조립한다',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
사용 예:
  # 기본 사용법
  # /path/to/base/ 아래에서 MyPresentation/ 폴더를 찾는다
  # MyPresentation/origin_image/ 에서 이미지를 읽는다
  # PPT를 MyPresentation/MyPresentation.pptx 로 저장한다
  python assemble_ppt.py /path/to/base/ MyPresentation.pptx

  # 4:3 가로세로 비율 지정
  python assemble_ppt.py /path/to/base/ MyPresentation.pptx --aspect-ratio 4:3

  # PPT 생성 없이 디렉터리만 초기화
  python assemble_ppt.py /path/to/base/ MyPresentation.pptx --init

폴더 구조 요구사항:
  /path/to/base/MyPresentation/
  ├── origin_image/
  │   ├── slide_01.png
  │   ├── slide_02.png
  │   └── ...
  ├── speech.md (선택, PPT 노트에 기록됨)
  └── MyPresentation.pptx (여기에 생성됨)

주의:
  - 이미지 파일은 반드시 origin_image 하위 디렉터리에 있어야 한다
  - slide_01.png, slide_02.png 형식의 정식 이미지만 읽으며 그 외 이미지는 무시된다
  - 이미지 파일은 페이지 번호순으로 정렬된다
  - 권장 파일명: slide_01.png, slide_02.png, ...
  - 각 이미지는 슬라이드 페이지 전체를 채운다
  - 프로젝트 디렉터리에 speech.md가 있으면 Slide N 제목 기준으로 페이지별 노트를 기록한다
        '''
    )

    parser.add_argument('base_dir', help='기준 디렉터리 (PPT 프로젝트 폴더의 상위 폴더)')
    parser.add_argument('output', help='출력 PPT 파일 이름 (.pptx)')
    parser.add_argument('--aspect-ratio', '--ar',
                        choices=['16:9', '4:3'],
                        default='16:9',
                        help='슬라이드 가로세로 비율 (기본값: 16:9)')
    parser.add_argument('--init',
                        action='store_true',
                        help='PPT 프로젝트 디렉터리와 origin_image 하위 디렉터리만 생성하고 PPT는 만들지 않는다')

    args = parser.parse_args()

    # 출력 파일에 .pptx 확장자 보장
    output_filename = args.output
    if not output_filename.lower().endswith('.pptx'):
        output_filename += '.pptx'

    # PPT 이름 추출 (확장자 제외)
    ppt_name = os.path.splitext(os.path.basename(output_filename))[0]

    # PPT 프로젝트 디렉터리 구성
    ppt_project_dir = os.path.join(args.base_dir, ppt_name)
    origin_image_dir = os.path.join(ppt_project_dir, "origin_image")

    if args.init:
        os.makedirs(origin_image_dir, exist_ok=True)
        print(f"✓ PPT 프로젝트 디렉터리 준비 완료: {ppt_project_dir}")
        print(f"✓ 슬라이드 이미지 디렉터리 준비 완료: {origin_image_dir}")
        sys.exit(0)

    if not os.path.exists(ppt_project_dir):
        print(f"오류: PPT 프로젝트 디렉터리가 없습니다: {ppt_project_dir}")
        print("디렉터리를 초기화하려면 --init 옵션을 추가하세요")
        sys.exit(1)

    if not os.path.exists(origin_image_dir):
        print(f"오류: origin_image 디렉터리가 없습니다: {origin_image_dir}")
        print("디렉터리를 초기화하려면 --init 옵션을 추가하세요")
        sys.exit(1)

    # 출력 경로 설정
    output_path = os.path.join(ppt_project_dir, output_filename)

    # 슬라이드 이미지 수집
    print(f"PPT 프로젝트 디렉터리 스캔 중: {ppt_project_dir}")
    image_files = get_slide_images(ppt_project_dir)

    if not image_files:
        print("오류: 이미지 파일을 찾을 수 없습니다")
        print("지원 형식: .png, .jpg, .jpeg, .gif, .bmp")
        print(f"\n슬라이드 이미지를 다음 위치에 넣으세요: {origin_image_dir}/")
        sys.exit(1)

    print(f"슬라이드 이미지 {len(image_files)}개 발견\n")
    speaker_notes = load_speaker_notes(ppt_project_dir)
    if speaker_notes:
        print(f"노트 {len(speaker_notes)}페이지 발견: {os.path.join(ppt_project_dir, 'speech.md')}\n")

    # 프레젠테이션 생성
    print(f"PPT 생성 중 (가로세로 비율: {args.aspect_ratio})...")
    print("-" * 50)

    success = create_presentation(image_files, output_path, args.aspect_ratio, speaker_notes)

    sys.exit(0 if success else 1)


if __name__ == '__main__':
    main()
